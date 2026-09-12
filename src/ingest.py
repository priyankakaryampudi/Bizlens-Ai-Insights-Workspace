from pathlib import Path
import io
import pandas as pd

SUPPORTED = ['csv','xlsx','xls','docx','pdf','pptx','txt']

def read_tabular(name, data):
    suffix = Path(name).suffix.lower()
    if suffix == '.csv':
        return pd.read_csv(io.BytesIO(data))
    if suffix in ('.xlsx', '.xls'):
        return pd.read_excel(io.BytesIO(data))
    raise ValueError('Unsupported tabular file')

def _clean_cell(text):
    return ' '.join(str(text or '').replace('(cid:127)', '').replace('\uf0b7', '').split())


def _pdf_table_sentences(table_rows):
    """Turn a detected PDF table into clean, distinct sentences instead of one
    run-on blob, so a requirement like FR-02 doesn't get merged with FR-01/FR-03."""
    rows = [[_clean_cell(c) for c in r] for r in table_rows]
    rows = [r for r in rows if any(r)]
    if len(rows) < 2:
        return []
    header = rows[0]
    out = []
    for row in rows[1:]:
        pairs = [f'{h}: {v}' for h, v in zip(header, row) if h and v]
        if pairs:
            out.append('; '.join(pairs) + '.')
    return out


def _extract_pdf_text(data):
    try:
        import pdfplumber
        parts = []
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                tables = page.find_tables()
                body_page = page
                for t in tables:
                    body_page = body_page.outside_bbox(t.bbox)
                raw = _clean_cell(body_page.extract_text() or '')
                if raw:
                    # Terminate each line before joining so a heading with no
                    # punctuation ("Functional Requirements") can't fuse with
                    # the next line's content during sentence-splitting.
                    for line in (body_page.extract_text() or '').splitlines():
                        line = _clean_cell(line)
                        if line and line[-1] not in '.!?:':
                            line += '.'
                        if line:
                            parts.append(line)
                for t in tables:
                    parts.extend(_pdf_table_sentences(t.extract() or []))
        return '\n'.join(parts)
    except Exception:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return '\n'.join((p.extract_text() or '') for p in reader.pages)


def extract_text(name, data):
    suffix = Path(name).suffix.lower()
    if suffix == '.txt':
        return data.decode('utf-8', 'ignore')
    if suffix == '.docx':
        from docx import Document
        doc = Document(io.BytesIO(data)); parts = []
        for p in doc.paragraphs:
            t = p.text.strip()
            if not t:
                continue
            # Terminate each paragraph so an unpunctuated heading ("Open
            # questions") can't fuse with the next paragraph's bullet text
            # once newlines are collapsed during sentence-splitting.
            if t[-1] not in '.!?:':
                t += '.'
            parts.append(t)
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    # Terminate each row so sentence-splitting can't merge it with the next row.
                    parts.append(' | '.join(cells).rstrip('.') + '.')
        return '\n'.join(parts)
    if suffix == '.pdf':
        return _extract_pdf_text(data)
    if suffix == '.pptx':
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data)); parts = []
        for i, slide in enumerate(prs.slides, 1):
            for shape in slide.shapes:
                if hasattr(shape, 'text') and shape.text.strip():
                    t = shape.text.strip()
                    # Terminate each shape's text so an unpunctuated title
                    # ("Questions Management Wants Answered") can't fuse with
                    # the next shape's content during sentence-splitting.
                    if t[-1] not in '.!?':
                        t += '.'
                    parts.append(f'[Slide {i}] {t}')
        return '\n'.join(parts)
    raise ValueError('Unsupported document file')

def classify(name):
    return 'dataset' if Path(name).suffix.lower() in ('.csv','.xlsx','.xls') else 'document'
